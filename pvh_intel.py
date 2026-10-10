"""
PVH intel export -- notes, stops and queries from the shared Supabase project,
joined to the PVH records: who drives what, and who is associated with whom.

Usage:
    python pvh_intel.py                 sign in, pull, build everything, open the map
    python pvh_intel.py --offline       rebuild from the last pull (intel/snapshot.json)
    python pvh_intel.py --no-map        build the files only
    python pvh_intel.py --snapshot F --out DIR   (testing: any snapshot, any folder)

Signs in as you (the email is remembered in pvh_local_config.json as
"intel_email"; the password is asked every run and never stored). Reads only
what the app itself can read, with the same public key the app uses.

Writes to intel/ (gitignored -- it is personal information):
    PVH_Intel.xlsx      Who drives what, People, Driver-Vehicle, Associations,
                        Profiles, Links, Notes, Activity, Monthly, Officers,
                        Hour x Weekday
    PVH_Intel_Map.html  the relationship map, view only
    link_reviews.csv    your links and decisions -- kept across every rebuild
    gephi_nodes.csv / gephi_edges.csv   Gephi "Import spreadsheet" format
    PVH_Intel.graphml   Gephi, Cytoscape, yEd, Neo4j
    snapshot.json       the raw pull, for --offline reruns

Then the map opens in the browser from a small server on this PC only
(127.0.0.1), where links can be made and removed: a person drives a vehicle,
or two people are associated. Each change is saved to link_reviews.csv and
every file above is rebuilt. Close the window to stop it.

How links are found:
    On file        owner -> vehicle, from the PVH records
    Note mention   a note on one record names another: plate, vehicle or
                   operator licence number, "deck 211", or a full name
    Same encounter one officer touched a vehicle and a person within 15
                   minutes, with a Stop or Note on one of them, and no more
                   than 2 vehicles / 2 people in that span (desk lookups are
                   not encounters)
    My link        made by you in the map
A driver-vehicle pair is On file, Company driver (same company) or NOT ON FILE
until you Confirm, Follow up or Remove it; two people named together in a note
are Seen together.
"""

import os
import re
import sys
import json
import getpass
import argparse
import datetime
import urllib.request
import urllib.error
from collections import Counter, defaultdict
from zoneinfo import ZoneInfo

HERE = os.path.dirname(os.path.abspath(__file__))
CFG_PATH = os.path.join(HERE, "pvh_local_config.json")
DATA_PATH = os.path.join(HERE, "PVH_data.json")
TZ = ZoneInfo("America/Halifax")
ENCOUNTER_SECS = 15 * 60
TABLES = ("notes", "checks", "queries")


# ------------------------------------------------------------------ helpers
def fmt_when(dt):
    """10/09/26 @ 1535Hrs -- same as the app."""
    return dt.strftime("%m/%d/%y @ %H%MHrs") if dt else ""


def parse_ts(s):
    if not s:
        return None
    s = re.sub(r"(\.\d{6})\d+", r"\1", str(s)).replace("Z", "+00:00")
    try:
        return datetime.datetime.fromisoformat(s).astimezone(TZ)
    except ValueError:
        return None


def s_(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v)


def alnum(s):
    return re.sub(r"[^A-Z0-9]", "", s_(s).upper())


# Notes and record labels are typed by officers, so text that a spreadsheet
# would read as a formula is kept as text. In the workbook the cell is typed
# as text; in a CSV it gets a leading apostrophe, which Excel hides and which
# load_reviews() takes off again.
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def as_text(cells):
    """openpyxl stores any string starting with "=" as a live formula."""
    for c in cells:
        if isinstance(c.value, str) and c.value.startswith("="):
            c.data_type = "s"


def csv_cell(v):
    return "'" + v if isinstance(v, str) and v.startswith(FORMULA_START) else v


def csv_uncell(v):
    return v[1:] if isinstance(v, str) and v.startswith("'") and v[1:].startswith(FORMULA_START) else v


# ------------------------------------------------------------------ companies
# Business names in the source are typed by hand: "CITY WIDE TAXI LTD - TOUR
# OPERATOR", "CITYWIDE TAXI", "OVER SEAS TAXI INC", "SIVLERSTAR LIMOUSINE",
# "CITY WIDE TAXI \ C-MAC TRANSPORTATION". Each part is cut down to its core
# ("CITYWIDE", "OVERSEAS", "SILVERSTAR") and cores are compared loosely.
CO_NOISE = set("""TAXI TAXIS CAB CABS LTD LIMITED INC INCORPORATED CO COMPANY CORP THE AND TOUR TOURS
OPERATOR OPERATORS SEE NOTES NOTE CHECK FILE RE SOT DECEASED NOVA SCOTIA NS LIMO LIMOUSINE
TRANSPORT TRANSPORTATION""".split())


def co_cores(name):
    out = set()
    for part in re.split(r"\s+-\s*|\s*-\s+|[/\\;,(]", s_(name).upper()):
        words = re.sub(r"[^A-Z ]", "", part.replace("'", "").replace("&", " ")).split()
        core = "".join(w for w in words if w not in CO_NOISE)
        if core:
            out.add(core)
    return out


def co_match(a, b):
    from difflib import SequenceMatcher
    for x in co_cores(a):
        for y in co_cores(b):
            n = min(len(x), len(y))
            if x == y or (n >= 4 and (x.startswith(y) or y.startswith(x))) or \
               (n >= 6 and SequenceMatcher(None, x, y).ratio() >= 0.85):
                return True
    return False


# ------------------------------------------------------------------ record keys
# These mirror clKey / clKeyV / nkW in the app exactly; a mismatch would
# orphan every note on that record.
def key_op(o):
    s = s_(o.get("Master Number")).strip().upper()
    s = re.sub(r"^4A\b", "", s)
    s = re.sub(r"^0*4\s+", "", s)
    s = re.sub(r"\s+", "", s)
    return "o:" + (s or "#" + (s_(o.get("ID")) or "?"))


def key_veh(v):
    s = re.sub(r"\s+", "", s_(v.get("Licence No")).strip().upper())
    return "v:" + (s or "#" + (s_(v.get("Vehicle ID")) or "?"))


def key_owner(w):
    nm = alnum(s_(w.get("Owner Last Name")) + s_(w.get("Owner First Name")) + s_(w.get("Business Name")))
    return "w:" + s_(w.get("Owner ID")).strip() + "|" + nm


# ------------------------------------------------------------------ Supabase
def http(method, url, key, token=None, body=None):
    h = {"apikey": key, "Content-Type": "application/json"}
    if token:
        h["Authorization"] = "Bearer " + token
    req = urllib.request.Request(url, method=method, headers=h,
                                 data=json.dumps(body).encode() if body is not None else None)
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            return r.status, json.loads(r.read() or b"null")
    except urllib.error.HTTPError as e:
        try:
            return e.code, json.loads(e.read() or b"null")
        except ValueError:
            return e.code, None


def pull(cfg):
    url = s_(cfg.get("supabase_url")).strip().rstrip("/")
    key = s_(cfg.get("supabase_anon_key")).strip()
    if not url or not key:
        sys.exit("pvh_local_config.json has no supabase_url / supabase_anon_key.")
    email = s_(cfg.get("intel_email")).strip()
    typed = input("Supabase email" + (" [" + email + "]" if email else "") + ": ").strip()
    if typed:
        email = typed
    pw = getpass.getpass("Password (not shown, not saved): ")
    st, j = http("POST", url + "/auth/v1/token?grant_type=password", key, body={"email": email, "password": pw})
    pw = None
    if st != 200 or not j or not j.get("access_token"):
        sys.exit("Sign-in refused: " + s_((j or {}).get("error_description") or (j or {}).get("msg") or st))
    token = j["access_token"]
    if typed and typed != cfg.get("intel_email"):
        try:
            raw = json.load(open(CFG_PATH, encoding="utf-8"))
            raw["intel_email"] = typed
            json.dump(raw, open(CFG_PATH, "w", encoding="utf-8"), indent=2)
        except (OSError, ValueError):
            pass
    snap = {"pulled_at": datetime.datetime.now(TZ).isoformat(), "by": email}
    for t in TABLES:
        rows, off = [], 0
        while True:
            st, page = http("GET", url + "/rest/v1/" + t + "?select=*&order=created_at.asc&limit=1000&offset=" + str(off), key, token)
            if st == 404 and t == "queries":
                print("  queries table not found -- run supabase/003_queries.sql. Continuing without it.")
                break
            if st != 200 or not isinstance(page, list):
                sys.exit("Could not read " + t + ": HTTP " + str(st) + " " + s_(page))
            rows += page
            if len(page) < 1000:
                break
            off += 1000
        snap[t] = rows
        print("  " + t + ": " + str(len(rows)))
    return snap


# ------------------------------------------------------------------ model
class Model:
    def __init__(self, db):
        self.V, self.O, self.W = db["vehicles"], db["operators"], db.get("owners", [])
        self.built = db.get("built", "")
        self.nodes = {}            # id -> dict
        self.alias = {}            # rec_key -> node id (owner-operators fold into the operator)
        self.veh_owner = {}        # vehicle node -> owner node (on file)
        for i, o in enumerate(self.O):
            k = key_op(o)
            self.add(k, "Operator", (" ".join(x for x in (s_(o.get("First Name")), s_(o.get("Middle")), s_(o.get("Last Name"))) if x) or "?"),
                     licence=s_(o.get("Licence Number")), company=s_(o.get("Business Name")),
                     detail=s_(o.get("Operator Type")) + " operator" + (" (inactive/cancelled)" if o.get("In Active") is True or o.get("Cancelled") is True else ""))
        for w in self.W:
            nk = key_owner(w)
            name = " ".join(x for x in (s_(w.get("Owner First Name")), s_(w.get("Owner Last Name"))) if x) or s_(w.get("Business Name")) or "?"
            if w.get("_op") is not None and w["_op"] < len(self.O):
                ok = key_op(self.O[w["_op"]])
                self.alias[nk] = ok
                self.nodes[ok]["type"] = "Owner-operator"
            else:
                self.add(nk, "Owner", name, company=s_(w.get("Business Name")), detail="Owner")
        for v in self.V:
            k = key_veh(v)
            deck = s_(v.get("Deck No"))
            ow = v.get("_owner")
            owner = self.resolve(key_owner(self.W[ow])) if ow is not None and ow < len(self.W) else None
            if owner:
                self.veh_owner[k] = owner
            owner_name = self.nodes[owner]["label"] if owner in self.nodes else \
                (" ".join(x for x in (s_(v.get("Owner First Name")), s_(v.get("Owner Last Name"))) if x)
                 or s_(v.get("Business Name")))
            desc = (s_(v.get("Make Model")) + " " + s_(v.get("Vehicle Color"))).strip()
            # Deck, owner, plate, description -- the order a car is read in the field.
            lab = " · ".join(x for x in (("Deck " + deck) if deck else "No deck", owner_name,
                                         s_(v.get("Plate No")), desc) if x)
            self.add(k, "Vehicle", lab, licence=s_(v.get("Licence No")), plate=s_(v.get("Plate No")), deck=deck,
                     company=s_(v.get("Business Name")), detail=s_(v.get("Vehicle Type")),
                     owner_name=owner_name, desc=desc)
        self.build_patterns()

    def add(self, nid, typ, label, **kw):
        if nid in self.nodes:
            return self.nodes[nid]
        n = {"id": nid, "type": typ, "label": label, "licence": "", "plate": "", "deck": "", "company": "",
             "detail": "", "in_data": True, "notes": 0, "stops": 0, "queries": 0, "events": [], "officers": Counter()}
        n.update(kw)
        self.nodes[nid] = n
        return n

    def resolve(self, rec_key):
        return self.alias.get(rec_key, rec_key)

    def node_for(self, rec_key, rec_label="", rec_type=""):
        nid = self.resolve(rec_key)
        if nid not in self.nodes:
            typ = {"v": "Vehicle", "o": "Operator", "w": "Owner"}.get(rec_key[:1], (rec_type or "Record").title())
            self.add(nid, typ, rec_label or rec_key, in_data=False, detail="Not in the current PVH data")
        return self.nodes[nid]

    # --- mention patterns, built once from the records
    def build_patterns(self):
        pats = []   # (compiled regex, [node ids], method, confidence)

        def sep_rx(code):
            return r"[\s-]?".join(re.escape(c) for c in code)

        plates, lic, decks, names = defaultdict(list), defaultdict(list), defaultdict(list), defaultdict(list)
        for v in self.V:
            k = key_veh(v)
            if len(alnum(v.get("Plate No"))) >= 4:
                plates[alnum(v.get("Plate No"))].append(k)
            if len(alnum(v.get("Licence No"))) >= 3:
                lic[alnum(v.get("Licence No"))].append(k)
            if s_(v.get("Deck No")).strip():
                decks[s_(v.get("Deck No")).strip()].append(k)
        for o in self.O:
            k = key_op(o)
            if len(alnum(o.get("Licence Number"))) >= 3:
                lic[alnum(o.get("Licence Number"))].append(k)
            f, m, l = (s_(o.get(x)).strip() for x in ("First Name", "Middle", "Last Name"))
            if len(f) >= 2 and len(l) >= 2:
                names[(f + " " + l).upper()].append(k)
                if m:
                    names[(f + " " + m + " " + l).upper()].append(k)
                names[(l + ", " + f).upper()].append(k)
        for w in self.W:
            if w.get("_op") is not None:
                continue
            f, l = s_(w.get("Owner First Name")).strip(), s_(w.get("Owner Last Name")).strip()
            if len(f) >= 2 and len(l) >= 2:
                names[(f + " " + l).upper()].append(key_owner(w))
        for code, ks in plates.items():
            pats.append((re.compile(r"(?<![A-Z0-9])" + sep_rx(code) + r"(?![A-Z0-9])", re.I), ks, "plate"))
        for code, ks in lic.items():
            m = re.match(r"([A-Z]+)(\d+)$", code)
            rx = (re.escape(m.group(1)) + r"\s?-?" + m.group(2)) if m else re.escape(code)
            pats.append((re.compile(r"(?<![A-Z0-9])" + rx + r"(?![A-Z0-9])", re.I), ks, "licence no."))
        for nm, ks in names.items():
            rx = r"\s+".join(re.escape(p) for p in nm.split())
            pats.append((re.compile(r"(?<![A-Z])" + rx + r"(?![A-Z])", re.I), sorted(set(ks)), "name"))
        self.pats = pats
        self.decks = decks
        self.deck_rx = re.compile(r"\b(?:deck|dk|car|unit|light)\s*(?:no\.?|number|#)?\s*#?\s*(\d{1,4})\b|#\s?(\d{1,4})\b", re.I)

    def mentions(self, text, self_id):
        out, seen = [], set()
        for rx, ks, method in self.pats:
            for m in rx.finditer(text):
                ids = sorted({self.resolve(k) for k in ks} - {self_id})
                for nid in ids:
                    if (nid, method) in seen:
                        continue
                    seen.add((nid, method))
                    out.append({"id": nid, "method": method, "text": m.group(0),
                                "confidence": "exact" if len(ids) == 1 else "possible (" + str(len(ids)) + " share it)"})
        for m in self.deck_rx.finditer(text):
            d = m.group(1) or m.group(2)
            ids = sorted(set(self.decks.get(d, [])) - {self_id})
            for nid in ids:
                if (nid, "deck") in seen or any(x["id"] == nid for x in out):
                    continue
                seen.add((nid, "deck"))
                out.append({"id": nid, "method": "deck", "text": m.group(0),
                            "confidence": "likely" if len(ids) == 1 else "possible (" + str(len(ids)) + " share it)"})
        # One mention per record per note: "<full name>, TD0000" is one
        # sighting, found two ways. Keep the strongest confidence.
        rank = lambda c: 0 if c == "exact" else (1 if c == "likely" else 2)
        merged = {}
        for m in out:
            cur = merged.get(m["id"])
            if not cur:
                merged[m["id"]] = dict(m)
                continue
            cur["method"] += " + " + m["method"]
            cur["text"] += "”, “" + m["text"]
            if rank(m["confidence"]) < rank(cur["confidence"]):
                cur["confidence"] = m["confidence"]
        return list(merged.values())


# ------------------------------------------------------------------ analysis
PERSON = ("Operator", "Owner-operator", "Owner")


def pair_id(a, b):
    """Person | vehicle for a drives pair; the two people in id order otherwise."""
    if b["type"] == "Vehicle":
        return a["id"] + " | " + b["id"]
    if a["type"] == "Vehicle":
        return b["id"] + " | " + a["id"]
    x, y = sorted((a["id"], b["id"]))
    return x + " | " + y


def new_pair(model, a, b):
    """A drives pair (a person and a vehicle) or an assoc pair (two people)."""
    p = {"id": pair_id(a, b), "links": [], "evidence": [], "review": None, "added": False, "reopened_by": []}
    if "Vehicle" in (a["type"], b["type"]):
        veh, per = (a, b) if a["type"] == "Vehicle" else (b, a)
        owner = model.veh_owner.get(veh["id"])
        if owner == per["id"]:
            base = "On file"
        elif co_match(per["company"], veh["company"]) or \
                (owner in model.nodes and co_match(per["company"], model.nodes[owner]["company"])):
            base = "Company driver"
        else:
            base = "NOT ON FILE"
        p.update(kind="drives", per=per, b=veh, veh=veh, owner=owner, base=base, status=base)
    else:
        x, y = sorted((a, b), key=lambda n: n["id"])
        p.update(kind="assoc", per=x, b=y, veh=None, owner=None, base="Seen together", status="Seen together")
    return p


def analyse(model, snap):
    events, notes, links = [], [], {}
    officers = {}

    def officer(row):
        oid = s_(row.get("author_id")) or ("name:" + s_(row.get("author_name")))
        officers.setdefault(oid, s_(row.get("author_name")) or "?")
        return oid

    def link(a, b, rel, when, ev):
        if a == b:
            return
        a, b = sorted((a, b))
        L = links.setdefault((a, b, rel), {"source": a, "target": b, "relation": rel, "count": 0,
                                          "first": None, "last": None, "officers": Counter(), "evidence": []})
        L["count"] += 1
        if when:
            L["first"] = min(L["first"] or when, when)
            L["last"] = max(L["last"] or when, when)
        if ev.get("officer"):
            L["officers"][ev["officer"]] += 1
        if len(L["evidence"]) < 20:
            L["evidence"].append(ev)

    for kind, table in (("Stop", "checks"), ("Query", "queries")):
        for r in snap.get(table, []):
            n = model.node_for(s_(r.get("rec_key")), s_(r.get("rec_label")))
            when = parse_ts(r.get("written_at"))
            oid = officer(r)
            n["stops" if kind == "Stop" else "queries"] += 1
            n["officers"][officers[oid]] += 1
            events.append({"kind": kind, "node": n["id"], "when": when, "oid": oid, "officer": officers[oid], "text": ""})

    for r in snap.get("notes", []):
        n = model.node_for(s_(r.get("rec_key")), s_(r.get("rec_label")), s_(r.get("rec_type")))
        when = parse_ts(r.get("written_at"))
        oid = officer(r)
        removed = bool(r.get("removed_at"))
        body = s_(r.get("body"))
        ms = [] if removed else model.mentions(body, n["id"])
        status = ("Removed by " + s_(r.get("removed_by"))) if removed else \
                 ("Done by " + s_(r.get("done_by"))) if r.get("done_at") else ("Open" if r.get("action_needed") else "")
        notes.append({"node": n["id"], "when": when, "officer": officers[oid], "body": body,
                      "action": bool(r.get("action_needed")), "status": status, "removed": removed,
                      "edited": (fmt_when(parse_ts(r.get("edited_at"))) + " by " + s_(r.get("edited_by"))) if r.get("edited_at") else "",
                      "mentions": ms})
        if not removed:
            n["notes"] += 1
            n["officers"][officers[oid]] += 1
            events.append({"kind": "Note", "node": n["id"], "when": when, "oid": oid, "officer": officers[oid], "text": body})
        for m in ms:
            model.node_for(m["id"])
            link(n["id"], m["id"], "Note mention", when,
                 {"officer": officers[oid], "when": when, "how": m["method"] + " \u201c" + m["text"] + "\u201d, " + m["confidence"],
                  "text": body, "confidence": m["confidence"]})

    # Same encounter: one officer, a vehicle and a person within 15 minutes --
    # but only when it looks like a stop. Opening records at a desk (checking
    # this very workbook against the app, say) touches many records in a few
    # minutes and is not an encounter, so a span with more than two vehicles or
    # two people is skipped, and there has to be a Stop or a Note on one of the
    # pair. Two queries alone are a lookup, not a sighting.
    by_off = defaultdict(list)
    for e in events:
        if e["when"]:
            by_off[e["oid"]].append(e)
    pair_last = {}
    pad = datetime.timedelta(minutes=5)
    for oid, evs in by_off.items():
        evs.sort(key=lambda e: e["when"])
        for i, a in enumerate(evs):
            for b in evs[i + 1:]:
                if (b["when"] - a["when"]).total_seconds() > ENCOUNTER_SECS:
                    break
                ta, tb = model.nodes[a["node"]]["type"], model.nodes[b["node"]]["type"]
                if (ta == "Vehicle") == (tb == "Vehicle"):
                    continue
                span = [e for e in evs if a["when"] - pad <= e["when"] <= b["when"] + pad]
                vehs = {e["node"] for e in span if model.nodes[e["node"]]["type"] == "Vehicle"}
                if len(vehs) > 2 or len({e["node"] for e in span}) - len(vehs) > 2:
                    continue
                if not any(e["kind"] != "Query" and e["node"] in (a["node"], b["node"]) for e in span):
                    continue
                pk = (oid,) + tuple(sorted((a["node"], b["node"])))
                lw = pair_last.get(pk)
                if lw and (a["when"] - lw).total_seconds() <= ENCOUNTER_SECS:
                    continue
                pair_last[pk] = b["when"]
                link(a["node"], b["node"], "Same encounter", a["when"],
                     {"officer": a["officer"], "when": a["when"],
                      "how": a["kind"] + " + " + b["kind"] + " within " + str(int((b["when"] - a["when"]).total_seconds() // 60)) + " min",
                      "text": a["text"] or b["text"], "confidence": "likely"})

    for e in events:
        model.nodes[e["node"]]["events"].append(e)

    # Pairs: one per person and vehicle ("drives"), and one per two people
    # ("assoc"), however many ways they were linked. A drives pair gets the
    # on-file / company check.
    pairs = {}
    for L in links.values():
        a, b = model.nodes[L["source"]], model.nodes[L["target"]]
        if a["type"] in PERSON and b["type"] in PERSON:
            P = pairs.get(pair_id(a, b)) or new_pair(model, a, b)
        elif {a["type"], b["type"]} & {"Vehicle"} and (a["type"] in PERSON or b["type"] in PERSON):
            P = pairs.get(pair_id(a, b)) or new_pair(model, a, b)
        else:
            continue
        pairs[P["id"]] = P
        P["links"].append(L)
        L["pair"] = P["id"]
    for P in pairs.values():
        P["evidence"] = sorted((dict(ev, relation=L["relation"]) for L in P["links"] for ev in L["evidence"]),
                               key=lambda e: e.get("when") or datetime.datetime.min.replace(tzinfo=TZ))
    return events, notes, links, pairs, officers


# ------------------------------------------------------------------ your links and decisions
# Everything you decide is kept in intel/link_reviews.csv, one row per pair, so
# it survives every rebuild:
#   * a decision on a pair the evidence found -- Confirmed, Follow up, Removed;
#   * a link you made yourself (added = yes) -- a person drives a vehicle, or
#     two people are associated -- which needs no evidence at all.
# Decisions come from the map (links are made and removed there) or from the
# Decision / Comment columns of the Driver-Vehicle and Associations sheets; a
# hidden "_was" column holds what each row was written with, so only cells
# actually changed in Excel count as edits. The CSV can be edited directly too.
#
# Pairs are keyed by the same record keys the phone app's shared notes use
# (v:..., o:..., w:...), so this file can later be shared or loaded into the
# shared database as it is.
REVIEW_FILE = "link_reviews.csv"
REVIEW_COLS = ["link_id", "kind", "person", "other", "added", "decision", "comment", "reviewed",
               "ev_last", "ev_officers"]
DECISIONS = ["Confirmed", "Follow up", "Removed"]
OPEN = {"NOT ON FILE", "Follow up", "Reopened"}
YOURS = {"Confirmed", "My link"}
SEP = " || "


def norm_decision(s):
    """Today's words, and the ones the first version used (Authorized driver,
    False link), so earlier decisions carry over."""
    t = s_(s).strip().lower()
    if not t:
        return ""
    if t.startswith(("auth", "conf", "ok", "yes")):
        return "Confirmed"
    if t.startswith(("false", "not", "rem", "wrong", "no", "del")):
        return "Removed"
    return "Follow up"


def load_reviews(dir_):
    import csv
    p = os.path.join(dir_, REVIEW_FILE)
    if not os.path.exists(p):
        return {}
    with open(p, newline="", encoding="utf-8-sig") as f:
        rows = [{k: csv_uncell(v) for k, v in r.items() if k} for r in csv.DictReader(f)]
    out = {}
    for r in rows:
        lid = s_(r.get("link_id")).strip()
        if not lid:
            continue
        r["link_id"] = lid
        r["person"] = s_(r.get("person") or r.get("driver"))
        r["other"] = s_(r.get("other") or r.get("vehicle"))
        r["kind"] = s_(r.get("kind")) or ("drives" if " | v:" in lid else "assoc")
        r["added"] = "yes" if s_(r.get("added")).strip().lower() in ("yes", "y", "true", "1") else ""
        r["decision"] = norm_decision(r.get("decision"))
        out[lid] = r
    return out


def save_reviews(dir_, reviews):
    import csv
    import shutil
    p = os.path.join(dir_, REVIEW_FILE)
    if os.path.exists(p):
        shutil.copyfile(p, p + ".bak")
    with open(p, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=REVIEW_COLS, extrasaction="ignore")
        w.writeheader()
        for k in sorted(reviews):
            w.writerow({c: csv_cell(v) for c, v in reviews[k].items()})


def harvest_edits(xlsx_path):
    """Decision / Comment cells the user changed in the last workbook."""
    if not os.path.exists(xlsx_path):
        return {}
    from openpyxl import load_workbook
    try:
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    except Exception as exc:   # a damaged or foreign file must not stop the run
        print("  could not read decisions from " + os.path.basename(xlsx_path) + ": " + str(exc))
        return {}
    out = {}
    for name in ("Driver-Vehicle", "Associations"):
        if name not in wb.sheetnames:
            continue
        rows = wb[name].iter_rows(values_only=True)
        head = [s_(h) for h in next(rows, [])]
        try:
            ci, cd, cc, cw = (head.index(x) for x in ("Link ID", "Decision", "Comment", "_was"))
        except ValueError:
            continue
        for r in rows:
            lid = s_(r[ci]).strip() if ci < len(r) else ""
            if not lid:
                continue
            dec, com = norm_decision(r[cd]), s_(r[cc]).strip()
            # The row was written with "_was"; a workbook from an earlier version
            # says "Authorized driver" where this one says "Confirmed", so compare
            # in today's words.
            was_dec, _, was_com = s_(r[cw]).partition(SEP)
            if (dec, com) != (norm_decision(was_dec), was_com.strip()):
                out[lid] = (dec, com)
    wb.close()
    return out


def baseline(r, P):
    """What the evidence looked like when the decision was made -- so later
    evidence the decision did not see can reopen it."""
    ev = [e for e in P["evidence"] if e.get("when")]
    r["ev_last"] = max(e["when"] for e in ev).isoformat() if ev else ""
    r["ev_officers"] = ";".join(sorted({e.get("officer", "") for e in P["evidence"]}))


def today_str():
    return datetime.datetime.now(TZ).strftime("%m/%d/%y")


def apply_links(model, links, pairs, out_dir, xlsx_path):
    """Fold your decisions and your own links into the pairs. Returns how many
    workbook edits were read."""
    reviews = load_reviews(out_dir)
    edits = harvest_edits(xlsx_path)
    today = today_str()
    changed = 0
    for lid, (dec, com) in list(edits.items()):
        r = reviews.get(lid) or {"link_id": lid, "added": ""}
        if (dec, com) == (norm_decision(r.get("decision")), s_(r.get("comment")).strip()):
            del edits[lid]     # already on file -- not a new review
            continue
        changed += 1
        if not dec and not com and not r.get("added"):
            reviews.pop(lid, None)
            continue
        r.update({"decision": dec, "comment": com, "reviewed": today})
        P = pairs.get(lid)
        if P:
            r.update(kind=P["kind"], person=P["per"]["label"], other=P["b"]["label"])
            baseline(r, P)
        reviews[lid] = r
    # A decision typed straight into the CSV has no baseline yet; take it now.
    for lid, r in reviews.items():
        if lid in pairs and r.get("decision") and not r.get("ev_last") and not r.get("ev_officers"):
            baseline(r, pairs[lid])
            changed += 1
    if changed:
        save_reviews(out_dir, reviews)

    # Your own links: a pair with no evidence needed, drawn as "My link".
    for lid, r in reviews.items():
        if not r.get("added"):
            continue
        a_id, _, b_id = lid.partition(" | ")
        if not a_id or not b_id:
            continue
        a = model.node_for(a_id, r.get("person"))
        b = model.node_for(b_id, r.get("other"))
        if not ((a["type"] in PERSON and b["type"] in PERSON) or
                (a["type"] in PERSON and b["type"] == "Vehicle")):
            continue
        P = pairs.get(lid)
        if not P:
            P = pairs[lid] = new_pair(model, a, b)
        P["added"] = True
        if norm_decision(r.get("decision")) != "Removed":
            when = None
            try:
                when = datetime.datetime.strptime(s_(r.get("reviewed")), "%m/%d/%y").replace(tzinfo=TZ)
            except ValueError:
                pass
            src, dst = sorted((a["id"], b["id"]))
            links[(src, dst, "My link")] = {
                "source": src, "target": dst, "relation": "My link", "count": 1, "first": when, "last": when,
                "officers": Counter(), "pair": lid,
                "evidence": [{"officer": "(you)", "when": when, "how": "Linked by you",
                              "text": s_(r.get("comment")), "confidence": "yours"}]}

    # Status: your decision first, then your own link, then what the evidence says.
    # A Confirmed or Removed pair comes back as Reopened when evidence arrives
    # that the decision did not see: another officer, or a new note naming them.
    for lid, P in pairs.items():
        r = reviews.get(lid)
        P["review"] = r
        dec = norm_decision(r.get("decision")) if r else ""
        if dec == "Follow up":
            P["status"] = "Follow up"
        elif dec in ("Confirmed", "Removed"):
            last = parse_ts(r.get("ev_last"))
            seen = set(s_(r.get("ev_officers")).split(";"))
            fresh = [e for e in P["evidence"] if last and e.get("when") and e["when"] > last
                     and (e.get("officer") not in seen or e.get("relation") == "Note mention")]
            P["status"] = "Reopened" if fresh else dec
            P["reopened_by"] = fresh
        elif P["added"]:
            P["status"] = "My link"
        else:
            P["status"] = P["base"]
    return len(edits)


def apply_edit(op, state):
    """One change made in the map. Returns (error message or None, id of the pair changed)."""
    N = state["model"].nodes
    pairs = state["pairs"]
    reviews = load_reviews(state["out"])
    today = today_str()
    comment = s_(op.get("comment")).strip()[:500]
    what = op.get("op")
    if what == "add":
        a, b = s_(op.get("a")), s_(op.get("b"))
        if a not in N or b not in N or a == b:
            return "That record is not in the current data.", None
        ta, tb = N[a]["type"], N[b]["type"]
        if not ((ta in PERSON and tb in PERSON) or (ta in PERSON and tb == "Vehicle") or
                (tb in PERSON and ta == "Vehicle")):
            return "A link joins a person to a vehicle (drives), or two people (associated).", None
        P = pairs.get(pair_id(N[a], N[b])) or new_pair(state["model"], N[a], N[b])
        r = reviews.get(P["id"]) or {"link_id": P["id"]}
        r.update(kind=P["kind"], person=P["per"]["label"], other=P["b"]["label"], added="yes", reviewed=today)
        if norm_decision(r.get("decision")) == "Removed":
            r["decision"] = ""
        if comment:
            r["comment"] = comment
        reviews[P["id"]] = r
        focus = P["id"]
    elif what == "decide":
        lid = s_(op.get("id"))
        dec = norm_decision(op.get("decision"))
        P = pairs.get(lid)
        if not P and lid not in reviews:
            return "That link is no longer in the data.", None
        r = reviews.get(lid) or {"link_id": lid, "added": ""}
        if P:
            r.update(kind=P["kind"], person=P["per"]["label"], other=P["b"]["label"])
            baseline(r, P)
        r.update(decision=dec, comment=comment, reviewed=today)
        if not dec and not r.get("added") and not comment:
            reviews.pop(lid, None)
        else:
            reviews[lid] = r
        focus = lid
    else:
        return "Unknown change.", None
    save_reviews(state["out"], reviews)
    return None, focus


def active_set(model, links):
    """Records with any activity, everything they link to, and their on-file
    owner / vehicles -- the part of the registry worth drawing."""
    act = {nid for nid, n in model.nodes.items() if n["notes"] or n["stops"] or n["queries"]}
    for L in links.values():
        act |= {L["source"], L["target"]}
    ring = set(act)
    for v, o in model.veh_owner.items():
        if v in act or o in act:
            ring |= {v, o}
    return act, ring


# ------------------------------------------------------------------ outputs
def write_xlsx(path, model, events, notes, links, pairs, officers, act, meta):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E79")
    red = PatternFill("solid", fgColor="F8D7DA")
    lilac = PatternFill("solid", fgColor="EAE4F7")
    grey = PatternFill("solid", fgColor="EDEFF2")
    edit = PatternFill("solid", fgColor="FFF4C2")
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(title, cols, rows, widths=None, wrap_cols=()):
        ws = wb.create_sheet(title)
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append(r)
            as_text(ws[ws.max_row])
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(c, max(10, min(40, len(c) + 4)))
            if c in wrap_cols:
                for cell in ws[get_column_letter(i)][1:]:
                    cell.alignment = wrap
        return ws

    N = model.nodes
    lab = lambda nid: N[nid]["label"]

    ws = wb.active
    ws.title = "Read me"
    for line in [
        ["PVH intel export"], [],
        ["Pulled", meta["pulled"]], ["Pulled by", meta["by"]], ["PVH data built", model.built],
        ["Notes / stops / queries", "%d / %d / %d" % (meta["n_notes"], meta["n_stops"], meta["n_queries"])], [],
        ["Sheets"],
        ["Who drives what", "Every vehicle in deck order: owner, drivers you have confirmed or linked, and drivers seen with it that still need a look."],
        ["People", "Everyone linked to a vehicle or another person: what they own, what they drive, who they drive for, who drives for them, who they are associated with."],
        ["Driver-Vehicle", "One row per person and vehicle linked by notes, an encounter or by you. Review them here or in the map."],
        ["Associations", "One row per two people linked by a note or by you."],
        ["Profiles", "Every record with activity: counts, last activity, who, and what it links to."],
        ["Links", "Every link with its evidence."],
        ["Notes", "All notes, with the records each one mentions."],
        ["Activity", "One row per query / stop / note -- use for pivots."],
        ["Monthly", "Activity per record per month."],
        ["Officers", "Activity per officer."],
        ["Hour x Weekday", "When activity happens."], [],
        ["Vehicles read", "Deck · owner · plate · description."], [],
        ["How links are found"],
        ["On file", "Owner -> vehicle in the PVH records."],
        ["Note mention", "A note on one record names another (plate, licence no., deck, full name)."],
        ["Same encounter", "One officer touched a vehicle and a person within 15 minutes, with a Stop or Note on one of them. Spans with more than 2 vehicles or 2 people (desk lookups) are ignored."],
        ["My link", "A link you made yourself in the map."],
        ["Confidence", "exact = one match; likely = deck number or encounter; possible = several records share that name or number -- verify."], [],
        ["Status"],
        ["On file", "The person is the vehicle's registered owner."],
        ["Company driver", "Not the owner, but the person and the vehicle are with the same company. Usually routine."],
        ["NOT ON FILE", "No registration or company connection. An open lead until reviewed."],
        ["Seen together", "Two people linked by a note. Not reviewed yet."],
        ["My link", "You linked them yourself."],
        ["Confirmed", "You reviewed it and it is right."],
        ["Follow up", "You want to keep an eye on it. Stays open."],
        ["Removed", "You took the link away. It stays on file, greyed out here and hidden in the map."],
        ["Reopened", "You confirmed or removed it, but new evidence has come in since: another officer, or a new note. Look again."], [],
        ["Making and reviewing links"],
        ["In the map", "Run Build_PVH_Intel.bat: the map opens in your browser. Click a record, then 'Link to…' and the other record. Click a line to Confirm, Follow up, Remove or Restore it. Changes are saved at once."],
        ["In this workbook", "On Driver-Vehicle or Associations, pick a Decision (yellow column), add a Comment, save, close and run Build_PVH_Intel.bat again."],
        ["", "Everything you decide is kept in link_reviews.csv and survives every rebuild."], [],
        ["This file holds personal information. Keep it on this PC; do not email or share it."],
    ]:
        ws.append(line)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110

    now = datetime.datetime.now(TZ)
    green = PatternFill("solid", fgColor="E2F4E8")
    SEEN = {"NOT ON FILE", "Company driver", "Reopened", "Follow up", "Seen together"}

    def lic(n):
        return n["label"] + (" (" + n["licence"] + ")" if n["type"] != "Vehicle" and n["licence"] else "")

    def short(v):
        return " · ".join(x for x in (("Deck " + v["deck"]) if v["deck"] else "", v["plate"]) if x) or v["label"]

    def last_seen(nid):
        ws_ = [e["when"] for e in N[nid]["events"] if e["when"]]
        return fmt_when(max(ws_)) if ws_ else ""

    live = [P for P in pairs.values() if P["status"] != "Removed"]
    by_veh, by_per = defaultdict(list), defaultdict(list)
    for P in live:
        by_per[P["per"]["id"]].append(P)
        (by_veh if P["kind"] == "drives" else by_per)[P["b"]["id"]].append(P)
    owns = defaultdict(list)
    for vid, oid in model.veh_owner.items():
        owns[oid].append(N[vid])

    # --- Who drives what: every vehicle, in deck order
    def deck_key(n):
        d = s_(n["deck"])
        return (0, int(d), n["plate"]) if d.isdigit() else (1, 0, n["plate"] or n["label"])
    rows = []
    for v in sorted((n for n in N.values() if n["type"] == "Vehicle" and n["in_data"]), key=deck_key):
        ps = [P for P in by_veh.get(v["id"], []) if P["status"] != "On file"]
        on = N.get(model.veh_owner.get(v["id"]))
        drivers = "; ".join(sorted({lic(P["per"]) for P in ps if P["status"] in YOURS}))
        seen = "; ".join(sorted({lic(P["per"]) + " [" + P["status"] + "]" for P in ps if P["status"] in SEEN}))
        rows.append([v["deck"], v.get("owner_name", ""), on["licence"] if on and on["type"] == "Owner-operator" else "",
                     v["plate"], v.get("desc", ""), v["detail"], drivers, seen, last_seen(v["id"]), v["id"]])
    ws = sheet("Who drives what", ["Deck", "Owner", "Owner-operator licence", "Plate", "Vehicle", "Type",
                                   "Drivers (confirmed / your links)", "Seen with -- to review", "Last activity", "Key"],
               rows, {"Owner": 30, "Owner-operator licence": 14, "Vehicle": 30, "Drivers (confirmed / your links)": 44,
                      "Seen with -- to review": 50, "Last activity": 18, "Key": 16},
               wrap_cols=("Drivers (confirmed / your links)", "Seen with -- to review"))
    for row in ws.iter_rows(min_row=2):
        f = red if row[7].value else green if row[6].value else None
        if f:
            for c in row:
                c.fill = f

    # --- People: who owns, drives and is associated with what and whom
    people = set()
    for P in live:
        if P["kind"] == "assoc":
            people |= {P["per"]["id"], P["b"]["id"]}
        elif P["status"] != "On file":
            people.add(P["per"]["id"])
            if P["owner"] in N:
                people.add(P["owner"])
    rows = []
    for pid in sorted(people, key=lambda k: N[k]["label"]):
        n = N[pid]
        mine = by_per.get(pid, [])
        drives = [P for P in mine if P["kind"] == "drives" and P["per"]["id"] == pid and P["status"] != "On file"]
        drives_for = {N[P["owner"]]["label"] for P in drives if P["status"] in YOURS and P["owner"] in N and P["owner"] != pid}
        their = {lic(P["per"]) for v in owns.get(pid, []) for P in by_veh.get(v["id"], [])
                 if P["status"] in YOURS and P["per"]["id"] != pid}
        assoc = []
        for P in mine:
            if P["kind"] != "assoc":
                continue
            o = P["b"] if P["per"]["id"] == pid else P["per"]
            assoc.append(lic(o) + ("" if P["status"] in YOURS else " [" + P["status"] + "]"))
        rows.append([n["label"], n["type"], n["licence"], n["company"],
                     "; ".join(short(v) for v in owns.get(pid, [])),
                     "; ".join(sorted(short(P["b"]) for P in drives if P["status"] in YOURS)),
                     "; ".join(sorted(short(P["b"]) + " [" + P["status"] + "]" for P in drives if P["status"] in SEEN)),
                     "; ".join(sorted(drives_for)), "; ".join(sorted(their)), "; ".join(sorted(assoc)), pid])
    sheet("People", ["Person", "Type", "Licence", "Company", "Owns", "Drives (confirmed / your links)",
                     "Seen driving -- to review", "Drives for", "Their drivers", "Associated with", "Key"],
          rows, {"Person": 28, "Company": 24, "Owns": 30, "Drives (confirmed / your links)": 34,
                 "Seen driving -- to review": 40, "Drives for": 30, "Their drivers": 34, "Associated with": 40, "Key": 22},
          wrap_cols=("Owns", "Drives (confirmed / your links)", "Seen driving -- to review", "Drives for",
                     "Their drivers", "Associated with"))

    # --- Profiles
    adj = defaultdict(set)
    for L in links.values():
        adj[L["source"]].add(L["target"])
        adj[L["target"]].add(L["source"])
    nof = Counter()
    for P in pairs.values():
        if P["kind"] == "drives" and P["status"] in OPEN:
            nof[P["b"]["id"]] += 1
            nof[P["per"]["id"]] += 1
    prof = []
    for nid in sorted(act, key=lambda k: -(N[k]["notes"] + N[k]["stops"] + N[k]["queries"])):
        n = N[nid]
        evs = sorted((e for e in n["events"] if e["when"]), key=lambda e: e["when"])
        last = evs[-1] if evs else None
        hrs = Counter(e["when"].hour for e in evs)
        prof.append([n["type"], n["label"], n["licence"] or n["plate"], n["company"], n["notes"], n["stops"], n["queries"],
                     sum(1 for e in evs if (now - e["when"]).days < 30),
                     fmt_when(last["when"]) if last else "", (last["kind"] + " · " + last["officer"]) if last else "",
                     len(n["officers"]), ", ".join(o for o, _ in n["officers"].most_common(3)),
                     ("%02d00-%02d59" % (hrs.most_common(1)[0][0], hrs.most_common(1)[0][0])) if hrs else "",
                     nof[nid] or "", "; ".join(sorted(lab(x) for x in adj[nid]))[:500],
                     "" if n["in_data"] else "Not in current data", nid])
    sheet("Profiles", ["Type", "Record", "Licence / plate", "Company", "Notes", "Stops", "Queries", "Last 30 days",
                       "Last activity", "Last by", "Officers", "Top officers", "Busiest hour", "Open leads",
                       "Linked to", "Flag", "Key"], prof,
          {"Record": 34, "Company": 24, "Linked to": 60, "Last activity": 18, "Last by": 22, "Top officers": 30, "Key": 22})

    # --- Driver-Vehicle and Associations: the pairs, with your decisions
    order = {"Reopened": 0, "Follow up": 1, "NOT ON FILE": 2, "Seen together": 3, "Company driver": 4,
             "My link": 5, "Confirmed": 6, "On file": 7, "Removed": 8}
    rank = lambda c: 0 if c == "exact" else (1 if c == "likely" else 2)

    def pair_cells(P):
        r, ev = P["review"] or {}, P["evidence"]
        dec, com = norm_decision(r.get("decision")), s_(r.get("comment"))
        dated = [e["when"] for e in ev if e.get("when")]
        notes_ev = [e for e in ev if e["relation"] == "Note mention"]
        status = P["status"]
        if status == "Reopened":
            # Blank the decision so picking it again counts as a fresh review.
            status += " (was " + dec + ") — new: " + \
                "; ".join(sorted({e.get("officer", "") + " " + fmt_when(e.get("when")) for e in P["reopened_by"]}))
            dec = ""
        kinds = Counter(e["relation"] for e in ev)
        if P["added"]:
            kinds["My link"] += 1
        return dict(status=status, dec=dec, com=com, reviewed=s_(r.get("reviewed")),
                    evidence=", ".join(k + (" ×" + str(c) if c > 1 else "") for k, c in kinds.items()),
                    times=len(ev), first=fmt_when(min(dated)) if dated else "", last=fmt_when(max(dated)) if dated else "",
                    officers=", ".join(sorted({e.get("officer", "") for e in ev})),
                    conf=min((e.get("confidence", "") for e in ev), key=rank, default=""),
                    how=ev[-1].get("how", "") if ev else ("Linked by you" if P["added"] else ""),
                    note=(notes_ev[-1].get("text", "") if notes_ev else "")[:500])

    def review_sheet(title, cols, rows, widths):
        ws = sheet(title, cols, rows, widths, wrap_cols=("Latest note", "Comment"))
        ws.column_dimensions[get_column_letter(len(cols))].hidden = True
        for row in ws.iter_rows(min_row=2):
            st = row[0].value or ""
            f = red if (st in OPEN or st.startswith("Reopened")) else lilac if st == "Company driver" else \
                green if st in YOURS else grey if st in ("On file", "Removed") else None
            if f:
                for c in row:
                    c.fill = f
            row[1].fill = row[2].fill = edit
        if rows:
            dv_ = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True,
                                 showErrorMessage=False)
            ws.add_data_validation(dv_)
            dv_.add("B2:B" + str(len(rows) + 1))

    rows = []
    for P in sorted((P for P in pairs.values() if P["kind"] == "drives"),
                    key=lambda P: (order.get(P["status"], 9), -len(P["evidence"]))):
        c, veh, per = pair_cells(P), P["b"], P["per"]
        rows.append([c["status"], c["dec"], c["com"], c["reviewed"], per["label"], per["licence"], per["company"],
                     veh["label"], veh["plate"], veh["deck"], veh["company"],
                     N[P["owner"]]["label"] if P["owner"] in N else "",
                     c["evidence"], c["times"], c["first"], c["last"], c["officers"], c["conf"], c["how"], c["note"],
                     P["id"], c["dec"] + SEP + c["com"]])
    review_sheet("Driver-Vehicle",
                 ["Status", "Decision", "Comment", "Reviewed", "Driver / person", "Licence", "Driver company", "Vehicle",
                  "Plate", "Deck", "Vehicle company", "Registered owner", "Evidence", "Times", "First", "Last",
                  "Officers", "Confidence", "Latest how found", "Latest note", "Link ID", "_was"], rows,
                 {"Status": 22, "Decision": 14, "Comment": 36, "Driver / person": 26, "Driver company": 24,
                  "Vehicle": 44, "Vehicle company": 24, "Registered owner": 24, "Evidence": 26, "First": 18, "Last": 18,
                  "Latest how found": 36, "Latest note": 60, "Link ID": 34})

    rows = []
    for P in sorted((P for P in pairs.values() if P["kind"] == "assoc"),
                    key=lambda P: (order.get(P["status"], 9), -len(P["evidence"]))):
        c, a, b = pair_cells(P), P["per"], P["b"]
        rows.append([c["status"], c["dec"], c["com"], c["reviewed"], a["label"], a["licence"], a["company"],
                     b["label"], b["licence"], b["company"], c["evidence"], c["times"], c["first"], c["last"],
                     c["officers"], c["how"], c["note"], P["id"], c["dec"] + SEP + c["com"]])
    review_sheet("Associations",
                 ["Status", "Decision", "Comment", "Reviewed", "Person A", "Licence A", "Company A", "Person B",
                  "Licence B", "Company B", "Evidence", "Times", "First", "Last", "Officers", "Latest how found",
                  "Latest note", "Link ID", "_was"], rows,
                 {"Status": 22, "Decision": 14, "Comment": 36, "Person A": 26, "Company A": 24, "Person B": 26,
                  "Company B": 24, "Evidence": 26, "First": 18, "Last": 18, "Latest how found": 30,
                  "Latest note": 60, "Link ID": 34})

    rows = []
    for L in sorted(links.values(), key=lambda L: -L["count"]):
        for ev in L["evidence"]:
            rows.append([lab(L["source"]), N[L["source"]]["type"], lab(L["target"]), N[L["target"]]["type"], L["relation"],
                         fmt_when(ev.get("when")), ev.get("officer", ""), ev.get("confidence", ""), ev.get("how", ""),
                         ev.get("text", "")[:500]])
    sheet("Links", ["Record A", "Type A", "Record B", "Type B", "Link", "When", "Officer", "Confidence", "How found", "Note"],
          rows, {"Record A": 30, "Record B": 30, "When": 18, "How found": 36, "Note": 60}, wrap_cols=("Note",))

    rows = []
    for nt in sorted(notes, key=lambda x: x["when"] or now, reverse=True):
        n = N[nt["node"]]
        rows.append([fmt_when(nt["when"]), nt["officer"], n["type"], n["label"], nt["body"], "Yes" if nt["action"] else "",
                     nt["status"], nt["edited"], "; ".join(lab(m["id"]) + " (" + m["method"] + ")" for m in nt["mentions"])])
    sheet("Notes", ["Written", "Officer", "Record type", "Record", "Note", "Action needed", "Status", "Edited", "Mentions"],
          rows, {"Written": 18, "Record": 32, "Note": 70, "Status": 22, "Edited": 26, "Mentions": 40}, wrap_cols=("Note", "Mentions"))

    rows = []
    for e in sorted(events, key=lambda e: e["when"] or now, reverse=True):
        n, w = N[e["node"]], e["when"]
        rows.append([w.replace(tzinfo=None) if w else None, fmt_when(w), w.strftime("%A") if w else "", w.hour if w else None,
                     w.strftime("%Y-%m") if w else "", e["kind"], e["officer"], n["type"], n["label"], e["node"]])
    ws = sheet("Activity", ["Date/time", "Written", "Weekday", "Hour", "Month", "Kind", "Officer", "Record type", "Record", "Key"],
               rows, {"Date/time": 18, "Written": 18, "Record": 34, "Key": 22})
    for c in ws["A"][1:]:
        c.number_format = "yyyy-mm-dd hh:mm"

    months = sorted({e["when"].strftime("%Y-%m") for e in events if e["when"]})
    per = defaultdict(Counter)
    for e in events:
        if e["when"]:
            per[e["node"]][e["when"].strftime("%Y-%m")] += 1
    rows = [[N[k]["type"], N[k]["label"]] + [per[k][m] or None for m in months] + [sum(per[k].values())]
            for k in sorted(per, key=lambda k: -sum(per[k].values()))]
    ws = sheet("Monthly", ["Type", "Record"] + months + ["Total"], rows, {"Record": 34})
    if rows and months:
        rng = "C2:" + get_column_letter(2 + len(months)) + str(len(rows) + 1)
        ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="5B9BD5"))

    offs = defaultdict(lambda: {"Query": 0, "Stop": 0, "Note": 0, "recs": set(), "hrs": Counter(), "last": None})
    for e in events:
        o = offs[e["officer"]]
        o[e["kind"]] += 1
        o["recs"].add(e["node"])
        if e["when"]:
            o["hrs"][e["when"].hour] += 1
            o["last"] = max(o["last"] or e["when"], e["when"])
    rows = [[k, o["Query"], o["Stop"], o["Note"], len(o["recs"]),
             ("%02d00" % o["hrs"].most_common(1)[0][0]) if o["hrs"] else "", fmt_when(o["last"])]
            for k, o in sorted(offs.items(), key=lambda kv: -(kv[1]["Query"] + kv[1]["Stop"] + kv[1]["Note"]))]
    sheet("Officers", ["Officer", "Queries", "Stops", "Notes", "Records touched", "Busiest hour", "Last activity"], rows,
          {"Officer": 22, "Last activity": 18})

    days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
    grid = Counter((e["when"].weekday(), e["when"].hour) for e in events if e["when"])
    rows = [["%02d00" % h] + [grid[(d, h)] or None for d in range(7)] for h in range(24)]
    ws = sheet("Hour x Weekday", ["Hour"] + days, rows)
    ws.conditional_formatting.add("B2:H25", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="F4B183"))
    wb.save(path)


def graph_rows(model, links, ring, pairs, officers_too=True):
    N = model.nodes
    nodes = [N[k] for k in sorted(ring)]
    edges = []
    for v, o in model.veh_owner.items():
        if v in ring and o in ring:
            edges.append({"source": o, "target": v, "relation": "On file", "count": 1, "first": None, "last": None,
                          "status": "", "open": False, "officers": Counter(), "evidence": []})
    # A pair's links carry its status (company driver, open lead, yours...);
    # a link you removed is left out.
    for L in links.values():
        P = pairs.get(L.get("pair"))
        if P and P["status"] == "Removed":
            continue
        L["status"] = P["status"] if P else ""
        L["open"] = bool(P) and P["status"] in OPEN
        L["comment"] = s_((P["review"] or {}).get("comment")) if P else ""
        edges.append(L)
    off_nodes, off_edges = {}, []
    if officers_too:
        for k in ring:
            for (off, kind), c in Counter((e["officer"], e["kind"]) for e in N[k]["events"]).items():
                oid = "officer:" + off
                off_nodes[oid] = {"id": oid, "type": "Officer", "label": off, "licence": "", "plate": "", "deck": "",
                                  "company": "", "detail": "Officer", "in_data": True, "notes": 0, "stops": 0,
                                  "queries": 0, "events": [], "officers": Counter()}
                off_edges.append({"source": oid, "target": k, "relation": {"Query": "Queried", "Stop": "Stopped", "Note": "Noted"}[kind],
                                  "count": c, "first": None, "last": None, "status": "", "open": False,
                                  "officers": Counter(), "evidence": []})
    return nodes + list(off_nodes.values()), edges + off_edges


def write_gephi(dir_, model, links, ring, pairs):
    import csv
    nodes, edges = graph_rows(model, links, ring, pairs)
    with open(os.path.join(dir_, "gephi_nodes.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Id", "Label", "Category", "Licence", "Plate", "Deck", "Company", "Notes", "Stops", "Queries", "InCurrentData"])
        for n in nodes:
            w.writerow([csv_cell(v) for v in (n["id"], n["label"], n["type"], n["licence"], n["plate"], n["deck"],
                                              n["company"], n["notes"], n["stops"], n["queries"], n["in_data"])])
    with open(os.path.join(dir_, "gephi_edges.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        # Gephi reads "Type" as Directed/Undirected, so the link kind goes in Relation.
        w.writerow(["Source", "Target", "Type", "Weight", "Relation", "Status", "OpenLead", "First", "Last"])
        for e in edges:
            w.writerow([csv_cell(v) for v in (e["source"], e["target"], "Undirected", e["count"], e["relation"],
                                              e["status"], "yes" if e["open"] else "",
                                              e["first"].isoformat() if e["first"] else "",
                                              e["last"].isoformat() if e["last"] else "")])
    from xml.sax.saxutils import escape as x
    na = [("label", "string"), ("category", "string"), ("licence", "string"), ("plate", "string"), ("company", "string"),
          ("notes", "int"), ("stops", "int"), ("queries", "int")]
    ea = [("relation", "string"), ("weight", "double"), ("status", "string"), ("openlead", "string"),
          ("first", "string"), ("last", "string")]
    out = ['<?xml version="1.0" encoding="UTF-8"?>',
           '<graphml xmlns="http://graphml.graphdrawing.org/xmlns">']
    out += ['<key id="n_%s" for="node" attr.name="%s" attr.type="%s"/>' % (k, k, t) for k, t in na]
    out += ['<key id="e_%s" for="edge" attr.name="%s" attr.type="%s"/>' % (k, k, t) for k, t in ea]
    out.append('<graph id="PVH" edgedefault="undirected">')
    for n in nodes:
        vals = {"label": n["label"], "category": n["type"], "licence": n["licence"], "plate": n["plate"], "company": n["company"],
                "notes": n["notes"], "stops": n["stops"], "queries": n["queries"]}
        out.append('<node id="%s">' % x(n["id"], {'"': "&quot;"}) + "".join('<data key="n_%s">%s</data>' % (k, x(s_(vals[k]))) for k, _ in na) + "</node>")
    for i, e in enumerate(edges):
        vals = {"relation": e["relation"], "weight": e["count"],
                "status": e["status"], "openlead": "yes" if e["open"] else "",
                "first": e["first"].isoformat() if e["first"] else "", "last": e["last"].isoformat() if e["last"] else ""}
        out.append('<edge id="e%d" source="%s" target="%s">' % (i, x(e["source"], {'"': "&quot;"}), x(e["target"], {'"': "&quot;"})) +
                   "".join('<data key="e_%s">%s</data>' % (k, x(s_(vals[k]))) for k, _ in ea) + "</edge>")
    out += ["</graph>", "</graphml>"]
    open(os.path.join(dir_, "PVH_Intel.graphml"), "w", encoding="utf-8").write("\n".join(out))


# ------------------------------------------------------------------ the map
def canvas_label(n):
    """Two short lines for the drawing; the full label is in the side panel."""
    cut = lambda s, k: s if len(s) <= k else s[:k - 1] + "…"
    if n["type"] == "Vehicle":
        top = " · ".join(x for x in (("Deck " + n["deck"]) if n["deck"] else "No deck", n.get("owner_name", "")) if x)
        bottom = " · ".join(x for x in (n["plate"], n.get("desc", "")) if x)
        return cut(top, 30) + ("\n" + cut(bottom, 30) if bottom else "")
    return cut(n["label"], 30) + ("\n" + n["licence"] if n["licence"] else "")


def map_data(model, links, notes, ring, pairs, meta):
    """Everything the map draws. Every record is included so any of them can
    be found and linked; only the active part of the registry starts visible."""
    N = model.nodes
    by_node = defaultdict(list)
    for nt in notes:
        if not nt["removed"]:
            by_node[nt["node"]].append({"w": fmt_when(nt["when"]), "o": nt["officer"], "b": nt["body"]})
    show = set(ring)
    for P in pairs.values():
        if P["status"] != "Removed":
            show |= {P["per"]["id"], P["b"]["id"]}
    for v, o in model.veh_owner.items():
        if v in show:
            show.add(o)
    nodes = []
    for nid, n in N.items():
        rec = {"id": nid, "type": n["type"], "label": n["label"], "l": canvas_label(n), "lic": n["licence"],
               "plate": n["plate"], "co": n["company"], "deck": n["deck"], "own": n.get("owner_name", ""),
               "desc": n.get("desc", ""), "n": n["notes"], "s": n["stops"], "q": n["queries"], "in": n["in_data"],
               "show": nid in show}
        if n["events"] or by_node.get(nid):
            evs = sorted((e for e in n["events"] if e["when"]), key=lambda e: e["when"], reverse=True)
            rec["ev"] = [[e["kind"], fmt_when(e["when"]), e["officer"]] for e in evs[:40]]
            rec["notes"] = by_node.get(nid, [])[:40]
        nodes.append(rec)

    def ev_rows(evidence):
        return [[fmt_when(e.get("when")), e.get("officer", ""), e.get("how", ""), s_(e.get("text"))[:300]] for e in evidence]

    edges = [{"id": "own|" + v, "from": o, "to": v, "kind": "owns"} for v, o in model.veh_owner.items()]
    for pid, P in pairs.items():
        if P["status"] == "On file":   # the owner seen with their own car -- the "owns" line says it
            continue
        r = P["review"] or {}
        edges.append({"id": pid, "from": P["per"]["id"], "to": P["b"]["id"], "kind": P["kind"], "st": P["status"],
                      "open": P["status"] in OPEN, "removed": P["status"] == "Removed", "mine": P["added"],
                      "dec": norm_decision(r.get("decision")), "cm": s_(r.get("comment")), "c": len(P["evidence"]),
                      "ev": ev_rows(P["evidence"][::-1][:10]),
                      "reop": sorted({e.get("officer", "") + " " + fmt_when(e.get("when")) for e in P["reopened_by"]})})
    i = 0
    for L in links.values():
        if L.get("pair") or L["relation"] == "My link":
            continue
        edges.append({"id": "lnk|%d" % i, "from": L["source"], "to": L["target"], "kind": "other",
                      "rel": L["relation"], "c": L["count"], "ev": ev_rows(L["evidence"][:10])})
        i += 1
    officers = {}
    for nid in ring:
        for (off, kind), c in Counter((e["officer"], e["kind"]) for e in N[nid]["events"]).items():
            officers["officer:" + off] = off
            edges.append({"id": "off|%d" % i, "from": "officer:" + off, "to": nid, "kind": "officer",
                          "rel": {"Query": "Queried", "Stop": "Stopped", "Note": "Noted"}[kind], "c": c})
            i += 1
    for oid, off in officers.items():
        nodes.append({"id": oid, "type": "Officer", "label": off, "l": off, "lic": "", "plate": "", "co": "",
                      "deck": "", "own": "", "desc": "", "n": 0, "s": 0, "q": 0, "in": True, "show": True})
    return {"nodes": nodes, "edges": edges, "meta": meta}


# The graph library is built into the map file, so the map opens offline and
# nothing on a third-party server runs inside a page that holds the intel. The
# copy in vendor/ is the official 9.1.9 release (byte-identical on cdnjs and
# npm, matching cdnjs's published hash). Without it, the CDN copy is used with
# that hash pinned, so the browser refuses anything that is not that file.
VIS_FILE = os.path.join(HERE, "vendor", "vis-network-9.1.9.min.js")
VIS_CDN = "https://cdnjs.cloudflare.com/ajax/libs/vis-network/9.1.9/standalone/umd/vis-network.min.js"
VIS_SRI = "sha512-iTgTmIgxyA2YehKNVbzLJx4j9SnuC5ihtRrtxVkXH/9nF3vXBN5YeNQp+6wufBWKD3u+roHVNOvWBMufQnBbug=="
_VIS = None


def vis_script():
    global _VIS
    if _VIS is None:
        try:
            with open(VIS_FILE, encoding="utf-8") as f:
                lib = f.read()
            _VIS = "<script>\n" + re.sub(r"</(script)", r"<\\/\1", lib, flags=re.I) + "\n</script>"
        except OSError:
            print("  vendor/vis-network-9.1.9.min.js not found -- the map will load the library online.")
            _VIS = '<script src="%s" integrity="%s" crossorigin="anonymous"></script>' % (VIS_CDN, VIS_SRI)
    return _VIS


def render_map(data, token=None):
    """The map page. With a token it can make and remove links through the
    local server (serve()); without one -- the file in intel/ -- it only views."""
    head, rest = MAP_HTML.split("__VISLIB__")
    mid, rest = rest.split("__DATA__")
    mid2, tail = rest.split("__EDIT__")
    edit = json.dumps({"token": token}) if token else "null"
    return head + vis_script() + mid + json.dumps(data).replace("</", "<\\/") + mid2 + edit + tail


def write_map(path, data):
    open(path, "w", encoding="utf-8").write(render_map(data))


MAP_HTML = r"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>PVH Intel Map</title>
__VISLIB__
<style>
:root{--bg:#f4f6f9;--panel:#fff;--line:#d8dee6;--text:#1b2430;--dim:#5b6675;--accent:#0a62c6;--bad:#c42b2b}
@media (prefers-color-scheme:dark){:root{--bg:#10151c;--panel:#18202a;--line:#2b3542;--text:#e6ebf1;--dim:#93a0b0;--accent:#5faeff;--bad:#ff6363}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--text);font:14px/1.4 -apple-system,"Segoe UI",Roboto,sans-serif;height:100vh;display:flex;flex-direction:column}
header{display:flex;gap:10px;align-items:center;flex-wrap:wrap;padding:8px 12px;border-bottom:1px solid var(--line);background:var(--panel)}
header b{font-size:15px;letter-spacing:.04em}header small{color:var(--dim)}
input[type=search]{flex:0 1 300px;padding:6px 10px;border:1px solid var(--line);border-radius:8px;background:var(--bg);color:var(--text)}
label{display:inline-flex;gap:4px;align-items:center;color:var(--dim);cursor:pointer}
#banner{display:none;padding:8px 12px;background:#fff4c2;color:#4a3a00;border-bottom:1px solid var(--line)}
#banner.on{display:block}#banner.err{background:#f8d7da;color:#6b1717}
#wrap{flex:1;display:flex;min-height:0}#net{flex:1;min-width:0}
#side{width:420px;max-width:45vw;overflow:auto;border-left:1px solid var(--line);background:var(--panel);padding:12px}
#side h2{font-size:16px;margin:0 0 4px}.k{color:var(--dim);font-size:12px}
.sec{margin-top:14px;font-size:11px;font-weight:700;letter-spacing:.08em;text-transform:uppercase;color:var(--dim)}
.it{border-top:1px solid var(--line);padding:6px 0}.it .k{display:block}
.kv{display:grid;grid-template-columns:70px 1fr;gap:2px 8px;margin:6px 0}.kv span:nth-child(odd){color:var(--dim)}
a{color:var(--accent);cursor:pointer}
.chip{display:inline-block;font-size:11px;font-weight:700;border-radius:4px;padding:1px 6px;margin-left:4px;white-space:nowrap}
.s-ok{background:#dcf2e4;color:#14532d}.s-open{background:#f8d7da;color:#7a1c1c}.s-co{background:#eae4f7;color:#4b2a8a}
.s-grey{background:#e6e9ee;color:#4b5563}.s-assoc{background:#dbeafe;color:#1e3a8a}
button.b{font:inherit;padding:6px 10px;border-radius:8px;border:1px solid var(--line);background:var(--bg);color:var(--text);cursor:pointer;margin:6px 6px 0 0}
button.b.p{background:var(--accent);border-color:var(--accent);color:#fff}button.b.d{color:var(--bad)}
textarea{width:100%;min-height:56px;margin-top:8px;font:inherit;padding:6px;border:1px solid var(--line);border-radius:8px;background:var(--bg);color:var(--text)}
.lg{display:inline-block;width:22px;height:0;border-top:3px solid;vertical-align:middle;margin-right:4px}.lg.d{border-top-style:dashed}
.ro{margin-top:10px;padding:8px;border-radius:8px;background:var(--bg)}
</style></head><body>
<header><b>PVH INTEL MAP</b><small id="meta"></small>
<input type="search" id="q" placeholder="Find a name, deck, plate or licence, then Enter">
<span id="types"></span>
<label><input type="checkbox" id="onlyopen"> Only open leads</label>
<label><input type="checkbox" id="showrem"> Show removed</label>
</header>
<div id="banner"></div>
<div id="wrap"><div id="net"></div><div id="side"></div></div>
<script>
let D=__DATA__;
const EDIT=__EDIT__;
const COL={"Vehicle":"#f0b429","Operator":"#3d8bfd","Owner-operator":"#7b61ff","Owner":"#8a94a3","Officer":"#2fb67c","Record":"#bbbbbb"};
const SHAPE={"Vehicle":"box","Officer":"diamond"};
const PERSON=new Set(["Operator","Owner-operator","Owner"]);
const YOURS=new Set(["Confirmed","My link"]);
const esc=s=>String(s==null?"":s).replace(/[&<>"']/g,c=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;"}[c]));
const show={};Object.keys(COL).forEach(t=>show[t]=t!=="Officer");
const onlyopen=document.getElementById("onlyopen"), showrem=document.getElementById("showrem"), side_=document.getElementById("side");
let byId={}, EB={}, shown=new Set(), openN=new Set(), LINK=null;
function index(){
  byId={};D.nodes.forEach(n=>{byId[n.id]=n;if(n.show)shown.add(n.id);});
  EB={};openN=new Set();
  D.edges.forEach(e=>{EB[e.id]=e;if(e.open&&!e.removed){openN.add(e.from);openN.add(e.to);}});
}
index();
const m=D.meta;
document.getElementById("meta").textContent="pulled "+m.pulled+" · "+m.n_notes+" notes · "+m.n_stops+" stops · "+m.n_queries+" queries";
document.getElementById("types").innerHTML=Object.keys(COL).filter(t=>D.nodes.some(n=>n.type===t)).map(t=>
  '<label><input type="checkbox" data-t="'+esc(t)+'" '+(show[t]?"checked":"")+'><span class="lg" style="border-color:'+COL[t]+';width:10px;border-top-width:10px;border-radius:5px"></span>'+esc(t)+'</label>').join(" ");
const act=n=>(n.n||0)+(n.s||0)+(n.q||0);
function chip(st){
  const c=YOURS.has(st)?"s-ok":(st==="NOT ON FILE"||st==="Follow up"||st==="Reopened")?"s-open":st==="Company driver"?"s-co":st==="Seen together"?"s-assoc":"s-grey";
  return '<span class="chip '+c+'">'+esc(st)+'</span>';
}
function estyle(e){
  if(e.kind==="owns")return {c:"#9aa5b1",w:1.5,d:false};
  if(e.kind==="officer")return {c:"#2fb67c",w:1,d:[2,4]};
  if(e.kind==="other")return {c:"#d9a21b",w:1,d:[4,4]};
  if(e.removed)return {c:"#c8ced6",w:1,d:[1,6]};
  const mine=YOURS.has(e.st);
  if(e.kind==="assoc")return {c:"#3d8bfd",w:mine?3:2,d:mine?false:[6,4]};
  if(mine)return {c:"#1f9d55",w:3,d:false};
  if(e.open)return {c:"#c42b2b",w:2.5,d:[8,5]};
  if(e.st==="Company driver")return {c:"#8e6fd1",w:2,d:[8,5]};
  return {c:"#9aa5b1",w:1.5,d:[8,5]};
}
const fg=()=>getComputedStyle(document.body).color;
function toNode(n){return {id:n.id,label:n.l||n.label,shape:SHAPE[n.type]||"dot",size:10+Math.min(30,act(n)*2),
  color:{background:COL[n.type]||"#bbb",border:n.in?"#0000":"#c42b2b"},borderWidth:n.in?0:2,
  font:{color:fg(),size:12},title:n.type+": "+n.label};}
function toEdge(e){const s=estyle(e);
  const what=e.kind==="owns"?"owns":e.kind==="drives"?"drives":e.kind==="assoc"?"associated":e.rel;
  return {id:e.id,from:e.from,to:e.to,width:s.w,dashes:s.d,color:{color:s.c,opacity:.9},
    label:(e.kind==="owns"||e.kind==="drives"||e.kind==="assoc")?what:undefined,font:{size:9,color:"#8a94a3",strokeWidth:0},
    title:what+(e.st?" — "+e.st:"")+(e.c>1?" ×"+e.c:"")};}
const nodes=new vis.DataSet(D.nodes.map(toNode));
const edges=new vis.DataSet(D.edges.map(toEdge));
const nodeView=new vis.DataView(nodes,{filter:n=>{const d=byId[n.id];
  return !!d&&show[d.type]&&shown.has(n.id)&&(!onlyopen.checked||openN.has(n.id));}});
const edgeView=new vis.DataView(edges,{filter:x=>{const e=EB[x.id];if(!e)return false;
  if(e.removed&&!showrem.checked)return false;if(e.kind==="officer"&&!show.Officer)return false;
  return !onlyopen.checked||e.open||e.kind==="owns";}});
const net=new vis.Network(document.getElementById("net"),{nodes:nodeView,edges:edgeView},
  {physics:{solver:"forceAtlas2Based",stabilization:{iterations:300}},interaction:{hover:true,multiselect:false}});
function refresh(){nodeView.refresh();edgeView.refresh();}
function banner(html,err){const b=document.getElementById("banner");b.className=html?("on"+(err?" err":"")):"";b.innerHTML=html||"";}
function reveal(id){shown.add(id);D.edges.forEach(e=>{if(e.kind!=="officer"&&(e.from===id||e.to===id)){shown.add(e.from);shown.add(e.to);}});}

/* --- side panel */
function home(){
  side_.innerHTML='<h2>Who drives what</h2><div class="k">Click a record or a line for details. Search finds any record in the PVH data, shown or not.</div>'+
    '<div class="sec">Lines</div><div class="it"><span class="lg" style="border-color:#9aa5b1"></span>owns (on file)<br>'+
    '<span class="lg" style="border-color:#1f9d55"></span>drives — confirmed or your link<br>'+
    '<span class="lg d" style="border-color:#c42b2b"></span>drives — not on file, follow up or reopened (open lead)<br>'+
    '<span class="lg d" style="border-color:#8e6fd1"></span>drives — company driver<br>'+
    '<span class="lg" style="border-color:#3d8bfd"></span>associated — confirmed or your link<br>'+
    '<span class="lg d" style="border-color:#3d8bfd"></span>associated — seen together in a note<br>'+
    '<span class="lg d" style="border-color:#d9a21b"></span>other note mention · <span class="lg d" style="border-color:#2fb67c"></span>officer activity (tick Officer)</div>'+
    (EDIT?'<div class="sec">Linking</div><div class="it">Click a person or vehicle, then <b>Link to…</b>, then the other record (or find it in the search box). Click a line to confirm, follow up, remove or restore it. Changes are saved at once.</div>'
         :'<div class="ro">This copy only views. To make or remove links, run <b>Build_PVH_Intel.bat</b>; the map opens in your browser ready for linking.</div>');
}
function item(id,extra){const n=byId[id];
  return '<div class="it"><a data-pick="'+esc(id)+'">'+esc(n?n.label:id)+'</a>'+(extra||"")+'</div>';}
function side(id){
  const n=byId[id];if(!n)return home();
  let h='<h2>'+esc(n.label)+'</h2><div class="k">'+esc(n.type)+(n.lic&&n.type!=="Vehicle"?' · '+esc(n.lic):'')+(n.co?' · '+esc(n.co):'')+
    (n.in?'':' · <span style="color:var(--bad)">not in current data</span>')+'</div>';
  if(n.type==="Vehicle")h+='<div class="kv"><span>Deck</span><span>'+esc(n.deck||"—")+'</span><span>Owner</span><span>'+esc(n.own||"—")+
    '</span><span>Plate</span><span>'+esc(n.plate||"—")+'</span><span>Vehicle</span><span>'+esc(n.desc||"—")+'</span><span>Licence</span><span>'+esc(n.lic||"—")+'</span></div>';
  if(n.type!=="Officer")h+='<div class="k">'+(n.n||0)+' notes · '+(n.s||0)+' stops · '+(n.q||0)+' queries</div>';
  if(EDIT&&n.type!=="Officer")h+='<button class="b p" data-act="linkfrom" data-id="'+esc(id)+'">Link to…</button>';
  const mine=D.edges.filter(e=>e.kind!=="officer"&&(e.from===id||e.to===id)&&(!e.removed||showrem.checked));
  const other=e=>e.from===id?e.to:e.from;
  const pair=e=>chip(e.st)+' <a class="k" data-edge="'+esc(e.id)+'">review</a>';
  const group=(title,list,fn)=>list.length?'<div class="sec">'+title+' ('+list.length+')</div>'+list.map(fn).join(""):"";
  if(n.type==="Vehicle"){
    h+=group("Owner",mine.filter(e=>e.kind==="owns"),e=>item(other(e)));
    h+=group("Drivers",mine.filter(e=>e.kind==="drives"),e=>item(other(e),pair(e)));
  }else{
    h+=group("Owns",mine.filter(e=>e.kind==="owns"),e=>item(other(e)));
    h+=group("Drives",mine.filter(e=>e.kind==="drives"),e=>item(other(e),pair(e)));
    h+=group("Associated with",mine.filter(e=>e.kind==="assoc"),e=>item(other(e),pair(e)));
  }
  h+=group("Other links",mine.filter(e=>e.kind==="other"),e=>item(other(e),' <span class="k">'+esc(e.rel)+'</span>'));
  if(n.notes&&n.notes.length)h+='<div class="sec">Notes</div>'+n.notes.map(x=>'<div class="it"><span class="k">'+esc(x.w)+' · '+esc(x.o)+'</span>'+esc(x.b)+'</div>').join("");
  if(n.ev&&n.ev.length)h+='<div class="sec">Activity</div>'+n.ev.map(x=>'<div class="it"><b>'+esc(x[0])+'</b> '+esc(x[1])+' · '+esc(x[2])+'</div>').join("");
  side_.innerHTML=h;
}
function edgeSide(eid){
  const e=EB[eid];if(!e)return home();
  const a=byId[e.from],b=byId[e.to];
  let h='<h2>'+(e.kind==="drives"?"Drives":"Associated")+' '+chip(e.st)+'</h2>'+item(e.from)+
    '<div class="k">'+(e.kind==="drives"?"drives":"associated with")+'</div>'+item(e.to);
  if(e.reop&&e.reop.length)h+='<div class="ro">New since your decision: '+esc(e.reop.join("; "))+'</div>';
  if(EDIT){
    h+='<textarea id="cmt" placeholder="Comment (optional)">'+esc(e.cm||"")+'</textarea>';
    const btn=(d,t,cls)=>'<button class="b '+(cls||"")+'" data-act="decide" data-id="'+esc(e.id)+'" data-dec="'+esc(d)+'">'+t+'</button>';
    if(e.removed)h+=btn("","Restore link","p");
    else{
      if(!YOURS.has(e.st))h+=btn("Confirmed","Confirm","p");
      if(e.st!=="Follow up")h+=btn("Follow up","Follow up");
      h+=btn("Removed","Remove link","d");
      if(e.dec&&e.st!=="Reopened")h+=btn("","Clear decision");
      h+=btn(e.dec||"","Save comment");
    }
  }else if(e.cm)h+='<div class="ro">Your comment: '+esc(e.cm)+'</div>';
  if(e.ev&&e.ev.length)h+='<div class="sec">Evidence ('+e.c+')</div>'+e.ev.map(v=>'<div class="it"><span class="k">'+esc(v[0])+' · '+esc(v[1])+' · '+esc(v[2])+'</span>'+esc(v[3])+'</div>').join("");
  else if(e.mine)h+='<div class="sec">Evidence</div><div class="it k">Linked by you; no sightings yet.</div>';
  side_.innerHTML=h;
}

/* --- picking, searching and linking */
function pick(id){
  if(LINK&&LINK!==id)return linkTarget(id);
  if(!byId[id])return;
  if(!show[byId[id].type]){show[byId[id].type]=true;const c=document.querySelector('[data-t="'+byId[id].type+'"]');if(c)c.checked=true;}
  reveal(id);refresh();net.selectNodes([id]);net.focus(id,{scale:1.2,animation:true});side(id);
}
const norm=s=>String(s||"").toUpperCase().replace(/[^A-Z0-9]/g,"");
const wordsOf=s=>String(s||"").toLowerCase().split(/[^a-z0-9]+/).filter(Boolean);
function findAll(q){
  const qa=norm(q), qw=wordsOf(q), dm=/^\s*(?:deck\s*)?#?\s*(\d{1,4})\s*$/i.exec(q);
  const out=[];
  D.nodes.forEach(n=>{
    if(n.type==="Officer")return;
    let s=0;
    if(dm&&n.deck===dm[1])s=100;
    else if(qa.length>=3&&(norm(n.plate)===qa||norm(n.lic)===qa))s=90;
    else if(qw.length&&qw.join("").length>=2){const w=wordsOf(n.label);if(qw.every(x=>w.some(y=>y.startsWith(x))))s=60;}
    if(s)out.push([s,n]);
  });
  out.sort((a,b)=>b[0]-a[0]||a[1].label.localeCompare(b[1].label));
  return out.map(x=>x[1]);
}
function startLink(id){
  LINK=id;
  banner('Linking from <b>'+esc(byId[id].label)+'</b>: click the vehicle or person to link to, or find it in the search box. <a data-act="cancel">Cancel</a> (Esc)');
}
function cancelLink(){LINK=null;banner("");}
function linkTarget(id){
  const a=byId[LINK],b=byId[id];if(!a||!b)return;
  const ok=(PERSON.has(a.type)&&PERSON.has(b.type))||(PERSON.has(a.type)&&b.type==="Vehicle")||(a.type==="Vehicle"&&PERSON.has(b.type));
  if(!ok){banner('A link joins a person to a vehicle, or two people. Pick another record, or <a data-act="cancel">cancel</a>.',true);return;}
  reveal(id);refresh();
  const per=PERSON.has(a.type)?a:b, oth=per===a?b:a;
  side_.innerHTML='<h2>New link</h2>'+item(per.id)+'<div class="k">'+(oth.type==="Vehicle"?"drives":"associated with")+'</div>'+item(oth.id)+
    '<textarea id="cmt" placeholder="Comment (optional) — e.g. seen driving it, drives for this owner, same address"></textarea>'+
    '<button class="b p" data-act="savelink" data-a="'+esc(a.id)+'" data-b="'+esc(b.id)+'">Save link</button><button class="b" data-act="cancel">Cancel</button>';
}
async function send(op){
  if(!EDIT)return null;
  let r;
  try{r=await fetch("/api/edit",{method:"POST",headers:{"Content-Type":"application/json","X-PVH-Token":EDIT.token},body:JSON.stringify(op)});}
  catch(e){banner("The intel tool is no longer running, so nothing was saved. Run Build_PVH_Intel.bat again to keep linking.",true);return null;}
  const j=await r.json().catch(()=>({ok:false,error:"The intel tool answered HTTP "+r.status+"."}));
  if(!j.ok){banner(esc(j.error||"Not saved."),true);return null;}
  D=j.data;index();
  nodes.update(D.nodes.map(toNode));
  const keep=new Set(D.nodes.map(n=>n.id));nodes.remove(nodes.getIds().filter(i=>!keep.has(i)));
  edges.clear();edges.add(D.edges.map(toEdge));refresh();
  banner(j.note?esc(j.note):"Saved.",!!j.note);if(!j.note)setTimeout(()=>{if(!LINK)banner("");},2500);
  return j;
}

/* --- events */
document.getElementById("types").addEventListener("change",e=>{show[e.target.dataset.t]=e.target.checked;refresh();});
onlyopen.addEventListener("change",refresh);showrem.addEventListener("change",refresh);
document.getElementById("q").addEventListener("keydown",e=>{
  if(e.key!=="Enter")return;const s=e.target.value.trim();if(!s)return;
  const hits=findAll(s);
  if(!hits.length){banner("Nothing in the PVH data matches “"+esc(s)+"”.",true);return;}
  if(hits.length===1)return pick(hits[0].id);
  side_.innerHTML='<h2>'+hits.length+' matches</h2>'+(LINK?'<div class="k">Pick the record to link to.</div>':'')+
    hits.slice(0,60).map(n=>item(n.id,'<span class="k">'+esc(n.type)+(n.lic&&n.type!=="Vehicle"?' · '+esc(n.lic):'')+'</span>')).join("");
});
document.addEventListener("keydown",e=>{if(e.key==="Escape"&&LINK){cancelLink();home();}});
net.on("click",p=>{
  if(p.nodes.length){const id=p.nodes[0];if(LINK)return linkTarget(id);return side(id);}
  if(p.edges.length){const e=EB[p.edges[0]];if(!e)return;if(e.kind==="drives"||e.kind==="assoc")return edgeSide(e.id);return side(e.from);}
});
document.body.addEventListener("click",async ev=>{
  const p=ev.target.closest("[data-pick]");if(p)return pick(p.dataset.pick);
  const g=ev.target.closest("[data-edge]");if(g){net.selectEdges([g.dataset.edge]);return edgeSide(g.dataset.edge);}
  const a=ev.target.closest("[data-act]");if(!a)return;
  const what=a.dataset.act, c=document.getElementById("cmt"), cm=c?c.value:"";
  if(what==="linkfrom")return startLink(a.dataset.id);
  if(what==="cancel"){cancelLink();return home();}
  if(what==="savelink"){const j=await send({op:"add",a:a.dataset.a,b:a.dataset.b,comment:cm});if(j){cancelLink();edgeSide(j.focus);}return;}
  if(what==="decide"){const j=await send({op:"decide",id:a.dataset.id,decision:a.dataset.dec,comment:cm});if(j)edgeSide(j.focus);}
});
home();
</script></body></html>
"""


# ------------------------------------------------------------------ build and serve
def rebuild(state):
    """Build every output from the snapshot and your links. Returns a note for
    the map ("" when all went well)."""
    model = Model(state["db"])
    events, notes, links, pairs, officers = analyse(model, state["snap"])
    n_edits = apply_links(model, links, pairs, state["out"], state["xl"])
    act, ring = active_set(model, links)
    note = ""
    try:
        write_xlsx(state["xl"], model, events, notes, links, pairs, officers, act, state["meta"])
    except PermissionError:
        note = "PVH_Intel.xlsx is open in Excel, so it was not updated. Close it; it updates with your next change."
    write_gephi(state["out"], model, links, ring, pairs)
    data = map_data(model, links, notes, ring, pairs, state["meta"])
    write_map(os.path.join(state["out"], "PVH_Intel_Map.html"), data)
    state.update(model=model, pairs=pairs, links=links, act=act, data=data, n_edits=n_edits)
    return note


def serve(state, port=8765, browser=True):
    """Show the map in the browser with linking turned on, until the window is
    closed. Only this PC can reach it (127.0.0.1); every change must carry the
    one-time key that is built into the page, so no other website can make one."""
    import hmac
    import secrets
    import threading
    import webbrowser
    from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
    token = secrets.token_urlsafe(24)
    where = {}
    lock = threading.Lock()   # one change at a time

    class Handler(BaseHTTPRequestHandler):
        # Browsers open spare connections and leave them idle; drop those
        # rather than letting them hold a request open.
        timeout = 15

        def reply(self, code, body, ctype="application/json; charset=utf-8"):
            data = body.encode("utf-8")
            self.send_response(code)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Length", str(len(data)))
            self.send_header("Cache-Control", "no-store")
            self.send_header("X-Content-Type-Options", "nosniff")
            self.end_headers()
            self.wfile.write(data)

        def host_ok(self):
            return self.headers.get("Host", "") in ("127.0.0.1:%d" % where["port"], "localhost:%d" % where["port"])

        def do_GET(self):
            if not self.host_ok():
                return self.reply(403, "Forbidden", "text/plain")
            if self.path.split("?")[0] not in ("/", "/index.html"):
                return self.reply(404, "Not found", "text/plain")
            with lock:
                page = render_map(state["data"], token)
            self.reply(200, page, "text/html; charset=utf-8")

        def do_POST(self):
            if not (self.host_ok() and self.path == "/api/edit"
                    and hmac.compare_digest(self.headers.get("X-PVH-Token", ""), token)):
                return self.reply(403, json.dumps({"ok": False, "error": "Not allowed."}))
            try:
                size = int(self.headers.get("Content-Length") or 0)
                op = json.loads(self.rfile.read(size)) if 0 < size <= 65536 else None
            except ValueError:
                op = None
            if not isinstance(op, dict):
                return self.reply(400, json.dumps({"ok": False, "error": "That change could not be read."}))
            with lock:
                try:
                    err, focus = apply_edit(op, state)
                    if err:
                        body = json.dumps({"ok": False, "error": err})
                    else:
                        note = rebuild(state)
                        body = json.dumps({"ok": True, "data": state["data"], "note": note, "focus": focus})
                        print("  saved: " + ("link added" if op.get("op") == "add" else (op.get("decision") or "decision cleared")), flush=True)
                except Exception as exc:   # say so in the map rather than leave it waiting
                    print("  could not save: " + repr(exc), flush=True)
                    body = json.dumps({"ok": False, "error": "That change could not be saved: " + str(exc)})
            self.reply(200, body)

        def log_message(self, fmt, *args):
            pass

    try:
        httpd = ThreadingHTTPServer(("127.0.0.1", port), Handler)
    except OSError:
        httpd = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    httpd.daemon_threads = True
    where["port"] = httpd.server_address[1]
    url = "http://127.0.0.1:%d/" % where["port"]
    print("")
    print("The map is open in your browser: " + url)
    print("Links you make or remove there are saved at once (intel\\link_reviews.csv),")
    print("and the workbook and map files are rebuilt each time.")
    print("Close this window (or press Ctrl+C) when you are done.", flush=True)
    if browser:
        webbrowser.open(url)
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        pass
    httpd.server_close()


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser(description="PVH intel export")
    ap.add_argument("--offline", action="store_true", help="rebuild from the last pull")
    ap.add_argument("--snapshot", help="snapshot file to read instead of pulling")
    ap.add_argument("--out", default=os.path.join(HERE, "intel"), help="output folder")
    ap.add_argument("--data", default=DATA_PATH, help="PVH_data.json")
    ap.add_argument("--no-map", action="store_true", help="write the files only; do not open the map for linking")
    ap.add_argument("--port", type=int, default=8765, help="port for the linking map (this PC only)")
    ap.add_argument("--no-browser", action="store_true", help="serve the map but do not open a browser")
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    snap_path = a.snapshot or os.path.join(a.out, "snapshot.json")
    if a.offline or a.snapshot:
        if not os.path.exists(snap_path):
            sys.exit("No snapshot at " + snap_path + " -- run once without --offline.")
        snap = json.load(open(snap_path, encoding="utf-8"))
        print("Using " + snap_path + " (pulled " + s_(snap.get("pulled_at")) + ")")
    else:
        cfg = json.load(open(CFG_PATH, encoding="utf-8")) if os.path.exists(CFG_PATH) else {}
        print("Pulling from Supabase…")
        snap = pull(cfg)
        json.dump(snap, open(snap_path, "w", encoding="utf-8"))
    if not os.path.exists(a.data):
        sys.exit("No " + a.data + " -- run the app build first.")
    state = {"snap": snap, "db": json.load(open(a.data, encoding="utf-8")), "out": a.out,
             "xl": os.path.join(a.out, "PVH_Intel.xlsx"),
             "meta": {"pulled": fmt_when(parse_ts(snap.get("pulled_at"))), "by": s_(snap.get("by")),
                      "n_notes": len(snap.get("notes", [])), "n_stops": len(snap.get("checks", [])),
                      "n_queries": len(snap.get("queries", []))}}
    note = rebuild(state)
    if state["n_edits"]:
        print("Decisions read from the workbook: %d (saved to %s)" % (state["n_edits"], REVIEW_FILE))
    if note:
        print("  " + note)
    pairs = state["pairs"]
    st = Counter(P["status"] for P in pairs.values() if P["kind"] == "drives")
    sa = Counter(P["status"] for P in pairs.values() if P["kind"] == "assoc")
    print("Records with activity: %d | links: %d | driver-vehicle: %d | associations: %d" % (
        len(state["act"]), len(state["links"]), sum(st.values()), sum(sa.values())))
    print("  open leads: %d  (NOT ON FILE %d, follow up %d, reopened %d)" % (
        sum(st[k] for k in OPEN), st["NOT ON FILE"], st["Follow up"], st["Reopened"]))
    print("  company driver: %d | confirmed: %d | your links: %d | removed: %d | on file: %d" % (
        st["Company driver"], st["Confirmed"] + sa["Confirmed"], st["My link"] + sa["My link"],
        st["Removed"] + sa["Removed"], st["On file"]))
    print("Written to " + a.out)
    if not a.no_map:
        serve(state, a.port, not a.no_browser)


if __name__ == "__main__":
    main()
